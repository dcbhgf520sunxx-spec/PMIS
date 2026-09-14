from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
ROOT=Path(__file__).resolve().parent
chapters=json.loads((ROOT/'章节内容.json').read_text())
manifest=json.loads((ROOT/'截图清单.json').read_text())
# Only include observed menu entries, in the captured UI order.
bykey={x['key']:x for x in chapters}
ordered=[bykey[x['key']] | {'capture':x} for x in manifest]
assert ordered and ordered[0]['key']=='home'
assert len(ordered)==13 and {x['key'] for x in ordered}=={x['key'] for x in chapters}
for ch in ordered:
 assert (ROOT/ch['capture']['file']).is_file(), ch['key']
 assert ch['capture'].get('verified'), ch['key']
assert ordered[0]['capture'].get('logoutVisible'), '首页必须显示退出登录'
d=Document();s=d.sections[0];s.orientation=WD_ORIENT.LANDSCAPE
s.page_width=Inches(11.69);s.page_height=Inches(8.27)
s.top_margin=Inches(.48);s.bottom_margin=Inches(.48);s.left_margin=Inches(.60);s.right_margin=Inches(.60)
s.header_distance=Inches(.20);s.footer_distance=Inches(.20)
for st in ['Normal','Title','Heading 1','Heading 2','Caption']:
 z=d.styles[st];z.font.name='Arial Unicode MS';z.font.color.rgb=RGBColor(0,0,0)
 z.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
 z.paragraph_format.space_before=Pt(0);z.paragraph_format.space_after=Pt(3)
 z.paragraph_format.line_spacing=Pt(13)
 z.paragraph_format.widow_control=False
 z.font.size=Pt(10.5)
d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 1'].font.bold=True
d.styles['Title'].font.size=Pt(15);d.styles['Title'].font.bold=True
h=s.header.paragraphs[0];h.text='智能数管 SIDM 企业数字化交付与运维管理软件 V1.0  操作说明书'
for r in h.runs:r.font.size=Pt(8)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('浙江中南建设集团有限公司    第 ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
f.add_run(' 页').font.size=Pt(8)
for ix,ch in enumerate(ordered,1):
 if ix>1:d.add_page_break()
 p=d.add_paragraph(f'{ix} {ch["title"]}',style='Title' if ix==1 else 'Heading 1');p.paragraph_format.line_spacing=Pt(20)
 p=d.add_paragraph('菜单位置：'+ch['entry']);p.paragraph_format.space_after=Pt(4)
 for r in p.runs:r.font.size=Pt(9)
 path=ROOT/ch['capture']['file'];w,h=Image.open(path).size
 scale=min(10.45/w,4.75/h)
 p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
 p.paragraph_format.line_spacing=1
 p.add_run().add_picture(str(path),width=Inches(w*scale),height=Inches(h*scale))
 p=d.add_paragraph(f'图 {ix}  {ch["title"]}完整界面',style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in p.runs:r.font.size=Pt(8)
 p=d.add_paragraph(ch['intro']);p.paragraph_format.space_after=Pt(4)
 for j,step in enumerate(ch['steps'],1):
  p=d.add_paragraph(f'{j}. {step}');p.paragraph_format.space_after=Pt(2)
 out=ROOT/'SIDM-V1.0-操作说明书-完整菜单版.docx'
for el in d.styles.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
for el in d.element.xpath('.//w:pBdr'):
 el.getparent().remove(el)
d.save(out)
print(out)
print('chapters',len(ordered),'images',len(d.inline_shapes))
